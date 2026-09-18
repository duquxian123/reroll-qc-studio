"""项目数据：分镜表 + 视频资源池。

模型变了
--------
以前每个分镜固定绑定自己的候选视频。现在改成**一个共享资源池**：

    分镜表（想拍什么）  +  资源池（手头有哪些视频片段）
                            ↓ 点「抽卡」
                    随机抽取 → 成为这个分镜的候选
                            ↓ 质检
                    通过判定 / 排名 / 采纳

这样"抽卡"就等价于真实场景里调用云端生成 API——每次拿回来的片子不一样。
"""

from __future__ import annotations

import csv
import json
import re
from dataclasses import dataclass, field
from pathlib import Path

from . import config
from .models import Shot


@dataclass
class PoolClip:
    """资源池里的一个视频片段。文件名里已经标好了状态。"""

    file: str                    # 相对 data/ 的路径，如 pool/clip_07_noise-heavy.mp4
    name: str                    # 纯文件名
    defects: list[str] = field(default_factory=list)
    severity: str = "none"
    verdict: str = "accept"      # accept / reject（测试集标注）
    base: str = ""
    from_shot: str = ""
    note: str = ""

    @property
    def is_good(self) -> bool:
        return not self.defects


@dataclass
class Project:
    project_id: str
    name: str
    shots: list[Shot] = field(default_factory=list)
    source: str = "builtin"           # builtin / imported
    warnings: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "project_id": self.project_id,
            "name": self.name,
            "source": self.source,
            "warnings": self.warnings,
            "shots": [s.to_dict() for s in self.shots],
        }


# ---------------------------------------------------------------------------
# 资源池
# ---------------------------------------------------------------------------


def load_pool() -> list[PoolClip]:
    """读取 data/pool 下的所有片段及其标签。"""
    if not config.POOL_LABELS.exists():
        return []
    try:
        raw = json.loads(config.POOL_LABELS.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError):
        return []

    clips: list[PoolClip] = []
    for item in raw.get("clips", []):
        name = item.get("file", "")
        if not name or not (config.POOL_DIR / name).exists():
            continue
        clips.append(
            PoolClip(
                file=f"pool/{name}",
                name=name,
                defects=list(item.get("defects", [])),
                severity=item.get("severity", "none"),
                verdict=item.get("verdict", "accept"),
                base=item.get("base", ""),
                from_shot=item.get("from_shot", ""),
                note=item.get("note", ""),
            )
        )
    return clips


def pool_is_empty() -> bool:
    return not load_pool()


# ---------------------------------------------------------------------------
# 分镜表
# ---------------------------------------------------------------------------

# ---------------------------------------------------------------------------
# 分镜表的列定义
# ---------------------------------------------------------------------------
#
# 模板是给人看的，所以表头用中文；解析时中英文都认，老模板不用改。
# FIELD_ORDER 同时就是模板里的列顺序。

FIELD_LABELS: dict[str, str] = {
    "shot_id": "分镜号",
    "scene": "场景",
    "prompt": "提示词",
    "negative_prompt": "负面词",
    "duration_s": "时长(秒)",
    "aspect_ratio": "画幅",
    "candidates_per_shot": "候选数",
    "ref_image": "参考图",
    "expected_faces": "期望人脸数",
    "expected_hands": "期望手部数",
    "priority": "优先级",
    "notes": "备注",
}
FIELD_ORDER: list[str] = list(FIELD_LABELS)

# 模板/CSV 的表头（中文）。保留旧名字，避免别处 import 失败。
CSV_FIELDS: list[str] = [FIELD_LABELS[key] for key in FIELD_ORDER]

FIELD_ALIASES: dict[str, tuple[str, ...]] = {
    "shot_id": ("shot_id", "shot id", "id", "分镜号", "镜头号", "分镜编号", "编号"),
    "scene": ("scene", "场景", "场次"),
    "prompt": ("prompt", "提示词", "画面描述", "描述"),
    "negative_prompt": (
        "negative_prompt", "negative prompt", "负面词", "负向词", "反向提示词",
    ),
    "duration_s": ("duration_s", "duration", "时长", "时长(秒)", "时长/s", "秒数"),
    "aspect_ratio": ("aspect_ratio", "aspect ratio", "画幅", "画幅比例", "宽高比", "比例"),
    "candidates_per_shot": (
        "candidates_per_shot", "candidates", "候选数", "候选条数", "生成条数", "抽卡数量",
    ),
    "ref_image": ("ref_image", "ref image", "参考图", "参考图片", "首帧图"),
    "expected_faces": ("expected_faces", "期望人脸数", "人脸数", "期望人脸"),
    "expected_hands": ("expected_hands", "期望手部数", "手部数", "期望手"),
    "priority": ("priority", "优先级", "优先"),
    "notes": ("notes", "note", "备注", "说明"),
}

# 优先级：内部用 high/medium/low，表格里写「高/中/低」或英文都认
PRIORITY_LABELS = {"high": "高", "medium": "中", "low": "低"}
_PRIORITY_ALIASES = {
    "high": "high", "高": "high", "高优先": "high", "重要": "high",
    "medium": "medium", "mid": "medium", "中": "medium", "中优先": "medium", "一般": "medium",
    "low": "low", "低": "low", "低优先": "low",
}

DEFAULT_CANDIDATES_PER_SHOT = 4
MAX_CANDIDATES_PER_SHOT = 12

SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".xlsm"}


def _norm_header(value: object) -> str:
    """表头归一化：忽略大小写、空白、下划线和全角括号。

    这样「时长（秒）」「时长 (秒)」「duration_s」「DURATION_S」都能对上同一列。
    """
    text = str(value).strip().lower()
    for a, b in (("（", "("), ("）", ")"), ("／", "/"), ("：", ":"), ("\u3000", "")):
        text = text.replace(a, b)
    return re.sub(r"[\s_\-·.]+", "", text)


_ALIAS_LOOKUP: dict[str, str] = {}
for _key, _names in FIELD_ALIASES.items():
    for _name in _names:
        _ALIAS_LOOKUP[_norm_header(_name)] = _key


def parse_table(path: Path) -> tuple[list[Shot], list[str]]:
    """解析分镜表。支持 .csv / .xlsx / .xlsm，列名中英文都认。

    Excel 只读**第一个工作表**，其它工作表会被忽略（并给出提示）。
    """
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        rows, notes = _read_xlsx(path)
    elif suffix == ".csv":
        rows, notes = _read_csv(path), []
    else:
        raise ValueError(f"不支持的文件格式 {suffix}，请用 .csv 或 .xlsx")

    shots, warnings = _rows_to_shots(rows)
    return shots, notes + warnings


def _read_csv(path: Path) -> list[dict]:
    text = path.read_text(encoding="utf-8-sig")
    return list(csv.DictReader(text.splitlines()))


def _read_xlsx(path: Path) -> tuple[list[dict], list[str]]:
    """用 openpyxl 读**第一个**工作表。空行会被丢掉。"""
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        ws = wb.worksheets[0]
        sheet_title = ws.title
        rows = [r for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()

    warnings: list[str] = []
    if len(names) > 1:
        warnings.append(
            f"文件里有 {len(names)} 个工作表（{', '.join(names)}），"
            f"只读取了第一个「{sheet_title}」。"
        )

    # 丢掉整行为空的行
    rows = [r for r in rows if any(c is not None and str(c).strip() for c in r)]
    if not rows:
        raise ValueError(f"工作表「{sheet_title}」里没有内容。")

    header = [_cell_text(c) for c in rows[0]]
    out: list[dict] = []
    for row in rows[1:]:
        out.append({header[i]: row[i] for i in range(min(len(header), len(row)))})
    return out, warnings


def _cell_text(value) -> str:
    """把单元格值统一成字符串（Excel 里的数字会变成 float）。"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _rows_to_shots(rows: list[dict]) -> tuple[list[Shot], list[str]]:
    warnings: list[str] = []
    if not rows:
        raise ValueError("表格是空的。")

    # 表头 → 内部字段名。中英文都认，认不出来的列收集起来提醒一句。
    mapped: list[dict] = []
    unknown: list[str] = []
    for row in rows:
        item: dict = {}
        for key, value in row.items():
            if key is None:
                continue
            canonical = _ALIAS_LOOKUP.get(_norm_header(key))
            if canonical:
                item[canonical] = value
            else:
                text = str(key).strip()
                if text and text not in unknown:
                    unknown.append(text)
        mapped.append(item)

    if "prompt" not in mapped[0]:
        raise ValueError(
            "表格缺少必填列「提示词」（英文 prompt 也可以）。"
            "建议直接用界面上的「下载分镜模板」。"
        )
    if unknown:
        warnings.append(f"忽略了无法识别的列：{'、'.join(unknown)}。")

    shots: list[Shot] = []
    for idx, row in enumerate(mapped, start=2):   # 第 1 行是表头
        get = lambda key: _cell_text(row.get(key))    # noqa: E731
        prompt = get("prompt")
        if not prompt:
            warnings.append(f"第 {idx} 行没有提示词，已跳过。")
            continue
        shot_id = get("shot_id") or f"IMP{idx:03d}"
        shots.append(
            Shot(
                shot_id=shot_id,
                scene=get("scene"),
                prompt=prompt,
                negative_prompt=get("negative_prompt"),
                expected_duration_s=_to_float(get("duration_s"), 5.0),
                aspect_ratio=get("aspect_ratio") or "16:9",
                ref_image=get("ref_image") or None,
                expected_faces=_to_int(get("expected_faces"), None),
                expected_hands=_to_int(get("expected_hands"), None),
                priority=_to_priority(get("priority")),
                candidates_per_shot=_to_int(
                    get("candidates_per_shot"), None, 1, MAX_CANDIDATES_PER_SHOT
                ),
                notes=get("notes"),
            )
        )

    if not shots:
        raise ValueError("表格里没有有效分镜（提示词列全空？）。")
    return shots, warnings


def load_builtin() -> Project:
    """载入内置分镜表。**初始没有候选**——要点了「抽卡」才会有。"""
    warnings: list[str] = []
    if not config.SHOTS_TEMPLATE.exists():
        return Project(
            project_id="builtin",
            name="内置样例项目",
            shots=[],
            warnings=["缺少 data/shots_template.csv，请先运行 tools/make_dataset.py。"],
        )

    try:
        shots, warnings = parse_table(config.SHOTS_TEMPLATE)
    except ValueError as exc:
        return Project(
            project_id="builtin", name="内置样例项目", shots=[], warnings=[str(exc)]
        )

    if pool_is_empty():
        warnings.append(
            "资源池是空的。请先运行 tools/make_dataset.py 生成片段（或双击 start.bat 自动生成）。"
        )

    return Project(
        project_id="builtin",
        name="内置样例项目",
        shots=shots,
        source="builtin",
        warnings=warnings,
    )


def blank_project() -> Project:
    """空白项目：一个分镜都没有。「初始化」按钮回到的就是它。"""
    return Project(project_id="blank", name="空白项目", shots=[], source="blank")


def import_csv(path: Path) -> Project:
    """导入分镜表（CSV / Excel）。只解析分镜，不生成候选——候选要靠「抽卡」。"""
    shots, warnings = parse_table(path)
    return Project(
        project_id=f"imported-{path.stem}",
        name=f"导入项目 {path.name}",
        shots=shots,
        source="imported",
        warnings=warnings,
    )


def _to_float(value: str | None, default: float) -> float:
    try:
        return float(value)  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return default


def _to_int(
    value: str | None,
    default: int | None,
    lo: int | None = None,
    hi: int | None = None,
) -> int | None:
    try:
        number = int(float(value))  # type: ignore[arg-type]
    except (TypeError, ValueError):
        return default
    if lo is not None and number < lo:
        return default
    if hi is not None and number > hi:
        return default
    return number


def _to_priority(value: str | None) -> str:
    """「高 / 中 / 低」和 high / medium / low 都认，认不出来按 medium。"""
    if not value:
        return "medium"
    return _PRIORITY_ALIASES.get(_norm_header(value), "medium")
