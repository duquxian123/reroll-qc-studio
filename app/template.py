"""分镜表模板（Excel）的生成。

为什么单独一个模块
------------------
模板是这个工作台的"入口文档"——用户下载它、在 Excel 里填、再传回来。
所以它必须：

* 表头是中文，一眼能看懂每列要填什么（鼠标悬停还有单元格批注）
* 有样式：表头底色、冻结窗格、合理列宽、下拉选项、斑马纹
* 只有**一个**工作表——导入时只读第一个表，多出来的表反而让人困惑

列顺序由 app/dataset.py 的 FIELD_ORDER 决定，和解析逻辑共用一份定义，
不会出现"模板里的列解析不了"这种事。
"""

from __future__ import annotations

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .dataset import (
    DEFAULT_CANDIDATES_PER_SHOT,
    FIELD_LABELS,
    FIELD_ORDER,
    MAX_CANDIDATES_PER_SHOT,
    PRIORITY_LABELS,
)
from .models import Shot

SHEET_TITLE = "分镜表"

HEADER_FILL = PatternFill("solid", fgColor="1D4ED8")
HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑", size=10.5)
BAND_FILL = PatternFill("solid", fgColor="F5F8FD")
_EDGE = Side(style="thin", color="D6DCE5")
BORDER = Border(left=_EDGE, right=_EDGE, top=_EDGE, bottom=_EDGE)

# 列宽：提示词和备注要宽，编号/数字类窄一点
COLUMN_WIDTH: dict[str, int] = {
    "shot_id": 10,
    "scene": 11,
    "prompt": 44,
    "negative_prompt": 18,
    "duration_s": 10,
    "aspect_ratio": 9,
    "candidates_per_shot": 9,
    "ref_image": 16,
    "expected_faces": 12,
    "expected_hands": 12,
    "priority": 9,
    "notes": 34,
}

# 鼠标悬停在表头上就能看到的填写说明
COLUMN_HINTS: dict[str, str] = {
    "shot_id": "分镜编号，例如 S001。留空会自动编号。",
    "scene": "这一幕在片子里的位置，例如「开场」「高潮」。只是给你自己看的。",
    "prompt": "必填。这一镜要生成什么画面，写得越具体越好。",
    "negative_prompt": "不想要的东西，例如「低质量,变形」。质检里的「模糊/闪烁」诊断会建议往这里加词。",
    "duration_s": "期望时长（秒），默认 5。质检的「规格」项会拿它和实际时长对比。",
    "aspect_ratio": "画幅，例如 16:9 / 9:16 / 1:1。",
    "candidates_per_shot": "这个分镜抽卡时抽几条候选（1-12），默认 4。抽卡相当于调用云端生成 API。",
    "ref_image": "可选。图生视频用的参考图路径，没有就留空。",
    "expected_faces": "这一镜期望出现几张人脸，没有人物就填 0 或留空。",
    "expected_hands": "这一镜期望出现几只手，没有手就留空。",
    "priority": "高 / 中 / 低，方便排优先级。",
    "notes": "给自己看的备注，例如「这一镜容易崩手」。不参与质检。",
}

# 这些列的文字左对齐并自动换行
WRAP_FIELDS = {"scene", "prompt", "negative_prompt", "ref_image", "notes"}
# 这些列居中
CENTER_FIELDS = {
    "shot_id", "duration_s", "aspect_ratio", "candidates_per_shot",
    "expected_faces", "expected_hands", "priority",
}

# 多留空行：用户往下继续填时，下拉选项和边框都已经就位
EXTRA_ROWS = 60


def build_workbook(shots: list[Shot]) -> Workbook:
    """按内置分镜生成模板工作簿。"""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_TITLE
    ws.sheet_properties.tabColor = "1D4ED8"

    columns = {field: index + 1 for index, field in enumerate(FIELD_ORDER)}
    last_col = get_column_letter(len(FIELD_ORDER))

    # ---- 表头 ----
    for field, index in columns.items():
        cell = ws.cell(row=1, column=index, value=FIELD_LABELS[field])
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
        hint = COLUMN_HINTS.get(field)
        if hint:
            comment = Comment(hint, "Reroll Studio")
            comment.width = 260
            comment.height = 110
            cell.comment = comment
    ws.row_dimensions[1].height = 26

    # ---- 数据行 ----
    for row_index, shot in enumerate(shots, start=2):
        banded = row_index % 2 == 0
        for field, index in columns.items():
            cell = ws.cell(row=row_index, column=index, value=_value_for(shot, field))
            cell.font = BODY_FONT
            cell.border = BORDER
            if banded:
                cell.fill = BAND_FILL
            if field in WRAP_FIELDS:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            elif field in CENTER_FIELDS:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row_index].height = 30

    last_row = max(2, len(shots) + 1 + EXTRA_ROWS)

    # 空白行也铺上边框，看起来像一张完整的表
    for row_index in range(len(shots) + 2, last_row + 1):
        for index in range(1, len(FIELD_ORDER) + 1):
            cell = ws.cell(row=row_index, column=index)
            cell.border = BORDER
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="center")

    # ---- 列宽 / 冻结 / 筛选 ----
    for field, index in columns.items():
        ws.column_dimensions[get_column_letter(index)].width = COLUMN_WIDTH.get(field, 14)

    # 冻结前两列 + 表头：往右滚动时「分镜号 / 场景」始终可见
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = f"A1:{last_col}{len(shots) + 1}"

    _add_validations(ws, columns, last_row)
    return wb


def _value_for(shot: Shot, field: str):
    """模板单元格里放什么值。"""
    if field == "duration_s":
        return shot.expected_duration_s or 5
    if field == "aspect_ratio":
        return shot.aspect_ratio or "16:9"
    if field == "candidates_per_shot":
        return shot.candidates_per_shot or DEFAULT_CANDIDATES_PER_SHOT
    if field == "expected_faces":
        return "" if shot.expected_faces is None else shot.expected_faces
    if field == "expected_hands":
        return "" if shot.expected_hands is None else shot.expected_hands
    if field == "priority":
        return PRIORITY_LABELS.get(shot.priority, "中")
    if field == "ref_image":
        return shot.ref_image or ""
    return getattr(shot, field, "") or ""


def _add_validations(ws, columns: dict[str, int], last_row: int) -> None:
    """下拉框 / 数字范围，减少填错。"""
    def add(validation: DataValidation, field: str) -> None:
        column = get_column_letter(columns[field])
        ws.add_data_validation(validation)
        validation.add(f"{column}2:{column}{last_row}")

    add(
        DataValidation(
            type="list",
            formula1='"高,中,低"',
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="优先级只能填 高 / 中 / 低",
            error="请从下拉列表里选：高、中、低",
        ),
        "priority",
    )
    add(
        DataValidation(
            type="list",
            formula1='"16:9,9:16,1:1,4:3"',
            allow_blank=True,
            showErrorMessage=False,
        ),
        "aspect_ratio",
    )
    add(
        DataValidation(
            type="whole",
            operator="between",
            formula1=1,
            formula2=MAX_CANDIDATES_PER_SHOT,
            allow_blank=True,
            showErrorMessage=True,
            errorTitle="候选数超出范围",
            error=f"候选数请填 1-{MAX_CANDIDATES_PER_SHOT} 之间的整数",
        ),
        "candidates_per_shot",
    )
    for field in ("expected_faces", "expected_hands"):
        add(
            DataValidation(
                type="whole",
                operator="between",
                formula1=0,
                formula2=8,
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="数量超出范围",
                error="请填 0-8 之间的整数",
            ),
            field,
        )
